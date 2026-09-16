from __future__ import annotations

import hashlib
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

from ..models import OrderLine, ValidationIssue


def normalize_header(v: object) -> str:
    if v is None:
        return ""
    return str(v).replace(" ", "").replace("\n", "").strip()


def read_all_sheets(path: Path) -> dict[str, list[tuple]]:
    """xlsx/xlsm -> openpyxl. xls -> xlrd. 시트명 -> 행 리스트(튜플)로 반환."""
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        wb = openpyxl.load_workbook(path, data_only=True)
        result = {}
        for sn in wb.sheetnames:
            ws = wb[sn]
            result[sn] = [
                tuple(row) for row in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True)
            ]
        wb.close()
        return result
    if suffix == ".xls":
        import xlrd

        wb = xlrd.open_workbook(str(path))
        result = {}
        for sheet in wb.sheets():
            rows = []
            for r in range(sheet.nrows):
                rows.append(tuple(sheet.cell_value(r, c) for c in range(sheet.ncols)))
            result[sheet.name] = rows
        return result
    raise ValueError(f"지원하지 않는 파일 형식입니다: {path.suffix} ({path.name})")


def file_hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


@dataclass
class ParseResult:
    lines: list[OrderLine] = field(default_factory=list)
    issues: list[ValidationIssue] = field(default_factory=list)
    format_name: str = "unknown"
