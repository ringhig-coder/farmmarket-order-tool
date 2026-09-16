"""마스터 발주정보(💰팜마켓 발주정보.xlsx) 로더.

절대 하지 않는 것:
- 🔒계정정보 시트를 열지 않는다.
- 각 시트의 로그인 ID/PW가 있는 상단 영역(무무식탁 1~9행, 나는농부 1~6행)을 읽지 않는다.
- 상품 단가를 코드에 하드코딩하지 않는다 (전부 이 파일에서 매번 읽는다).
"""
from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

import openpyxl

from .config import SupplierRules
from .models import MasterProduct

_ANNOTATION_PATTERNS = [
    r"^NEW$",
    r"^\*.*",
    r"^\(.*\)$",
    r".*인상.*",
    r".*거래중단.*",
    r"^중단.*",
    r".*보류.*",
    r".*상시.*할인.*",
    r".*쿠팡\s*x.*",
    r".*유리병\s*중단.*",
    r".*파우치\s*대체.*",
    r".*월정산.*",
    r".*가격변동.*",
]
_ANNOTATION_RE = [re.compile(p, re.IGNORECASE) for p in _ANNOTATION_PATTERNS]


def _is_annotation_line(line: str) -> bool:
    line = line.strip()
    if not line:
        return True
    return any(rx.match(line) for rx in _ANNOTATION_RE)


def clean_company_name(raw: object) -> str:
    """'해남 고구마 식품\\n\\n*단가인상\\n(26.04.16)' -> '해남고구마식품' 처럼 정규화."""
    if raw is None:
        return ""
    lines = str(raw).split("\n")
    kept = [ln.strip() for ln in lines if not _is_annotation_line(ln)]
    joined = " ".join(kept).strip()
    return re.sub(r"\s+", "", joined)


def clean_product_name(raw: object) -> str:
    if raw is None:
        return ""
    text = str(raw).replace("\n", " ").strip()
    return re.sub(r"\s+", " ", text)


def _normalize_header(v: object) -> str:
    if v is None:
        return ""
    return str(v).replace(" ", "").replace("\n", "")


_HEADER_TARGETS = {
    "company": "업체명",
    "product": "제품명",
    "price_ex": "공급가(택전)",
    "price_incl": "공급가(택포)",
    "shipping_fee": "택배비",
    "account": "계좌번호",
    "note": "비고",
}

# 시트마다 로그인 정보가 차지하는 상단 영역이 달라서, 헤더 행을 직접 지정한다.
_SHEET_HEADER_ROW = {
    "🍽️무무식탁": 10,
    "🌳나는농부": 7,
}


def _build_column_map(ws, header_row: int) -> dict[str, int]:
    # 시트 앞쪽(B~G열 부근)에 고객 판매가 계산용 '택배비' 등 이름이 겹치는 컬럼이 있어서,
    # 먼저 '업체명' 컬럼을 찾고 그 이후 범위에서만 나머지 헤더를 찾는다.
    company_col = None
    for col in range(1, ws.max_column + 1):
        if _normalize_header(ws.cell(row=header_row, column=col).value) == "업체명":
            company_col = col
            break
    if company_col is None:
        raise ValueError(f"마스터 시트 헤더 행({header_row})에서 '업체명' 컬럼을 찾지 못했습니다.")

    col_map: dict[str, int] = {}
    search_start = max(1, company_col - 1)
    for col in range(search_start, ws.max_column + 1):
        header_text = _normalize_header(ws.cell(row=header_row, column=col).value)
        for key, target in _HEADER_TARGETS.items():
            if header_text == target and key not in col_map:
                col_map[key] = col
    missing = set(_HEADER_TARGETS) - set(col_map)
    if missing:
        raise ValueError(
            f"마스터 시트 헤더에서 다음 컬럼을 찾지 못했습니다: {missing} "
            f"(헤더 행={header_row}). 마스터 파일 양식이 바뀌었을 수 있습니다."
        )
    return col_map


def _to_float(v: object) -> float | None:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    text = str(v).strip()
    if not text:
        return None
    text = text.replace(",", "")
    try:
        return float(text)
    except ValueError:
        return None  # "모름", "착불" 등 숫자가 아닌 값 -> 계산 불가로 취급


@dataclass
class MasterCatalog:
    products: list[MasterProduct]

    def companies(self) -> set[str]:
        return {p.company for p in self.products}

    def find_exact(self, company: str, product_name: str) -> MasterProduct | None:
        for p in self.products:
            if p.company == company and p.product_name == product_name:
                return p
        return None

    def find_by_company(self, company: str) -> list[MasterProduct]:
        return [p for p in self.products if p.company == company]

    def find_all_by_product_name(self, product_name: str) -> list[MasterProduct]:
        return [p for p in self.products if p.product_name == product_name]

    def shipping_fee_for(self, company: str) -> float | None:
        """같은 업체의 사이즈별 상품 행 중 택배비가 채워진 첫 값을 그 업체의 배송비로 본다."""
        for p in self.products:
            if p.company == company and p.shipping_fee is not None:
                return p.shipping_fee
        return None

    def account_for(self, company: str) -> str | None:
        """계좌번호도 업체 그룹의 첫 상품 행에만 채워져 있는 경우가 많아 같은 방식으로 찾는다."""
        for p in self.products:
            if p.company == company and p.account is not None:
                return p.account
        return None


def load_master_catalog(master_path: Path, rules: SupplierRules) -> MasterCatalog:
    wb = openpyxl.load_workbook(master_path, data_only=True)
    products: list[MasterProduct] = []

    for sheet_name, header_row in _SHEET_HEADER_ROW.items():
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"마스터 파일에 '{sheet_name}' 시트가 없습니다. 파일이 바뀌었는지 확인하세요.")
        ws = wb[sheet_name]
        col_map = _build_column_map(ws, header_row)
        last_company = ""

        for row in range(header_row + 1, ws.max_row + 1):
            raw_company = ws.cell(row=row, column=col_map["company"]).value
            raw_product = ws.cell(row=row, column=col_map["product"]).value

            if raw_company not in (None, ""):
                last_company = rules.resolve_company(clean_company_name(raw_company))

            product_name = clean_product_name(raw_product)
            if not product_name or not last_company:
                continue

            account = ws.cell(row=row, column=col_map["account"]).value
            if account not in (None, ""):
                account = re.sub(r"\s+", " ", str(account)).strip()
            else:
                account = None
            override_account, _warning = rules.account_override(last_company, account)
            if override_account:
                account = override_account

            note = ws.cell(row=row, column=col_map["note"]).value
            note = str(note).strip() if note not in (None, "") else None

            products.append(
                MasterProduct(
                    sheet=sheet_name,
                    row=row,
                    company=last_company,
                    product_name=product_name,
                    supply_price_ex_shipping=_to_float(ws.cell(row=row, column=col_map["price_ex"]).value),
                    supply_price_incl_shipping=_to_float(ws.cell(row=row, column=col_map["price_incl"]).value),
                    shipping_fee=_to_float(ws.cell(row=row, column=col_map["shipping_fee"]).value),
                    account=account,
                    note=note,
                )
            )

    wb.close()
    return MasterCatalog(products=products)


SELF_SUPPLY_SHEET = "🌳나는농부"
