"""발주서 자동 생성 테스트: 쓴 파일을 기존 파서로 다시 읽었을 때 내용이 그대로 맞아야 한다."""
import datetime

from farmmarket.order_writer import GeneratedOrderLine, load_output_templates, write_supplier_order_file
from farmmarket.parsers.simple_option import parse as parse_simple_option


def test_write_and_reread_roundtrip(project_root, tmp_path):
    templates = load_output_templates(project_root / "config" / "output_templates.json")
    lines = [
        GeneratedOrderLine(
            recipient="홍길동",
            phone1="010-1234-5678",
            address="서울시 강남구 어딘가",
            note="문 앞에 놔주세요",
            product_name="프리미엄 넛츠양갱",
            quantity=1,
        )
    ]
    out_path = tmp_path / "밀양한천.xlsx"
    write_supplier_order_file("밀양한천", lines, templates, out_path, order_date=datetime.date(2026, 9, 16))

    import openpyxl

    wb = openpyxl.load_workbook(out_path)
    ws = wb["발주발송관리"]
    rows = list(ws.iter_rows(values_only=True))
    parsed = parse_simple_option(out_path, rows)
    assert len(parsed.lines) == 1
    line = parsed.lines[0]
    assert line.recipient == "홍길동"
    assert line.address == "서울시 강남구 어딘가"
    assert line.product_name_raw == "프리미엄 넛츠양갱"
    assert line.quantity == 1


def test_write_farmmarket_with_extra_fields(project_root, tmp_path):
    templates = load_output_templates(project_root / "config" / "output_templates.json")
    lines = [
        GeneratedOrderLine(
            recipient="최승훈",
            phone1="010-3770-9462",
            zipcode="33403",
            address="충청남도 보령시 천북면 세편길 48",
            note="배송 전 미리 연락해 주세요",
            product_name="충전식 전동 분무기 BMC 18L",
            quantity=2,
            channel="스토어",
        )
    ]
    out_path = tmp_path / "팜마켓.xlsx"
    write_supplier_order_file("팜마켓", lines, templates, out_path)

    import openpyxl

    wb = openpyxl.load_workbook(out_path)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    parsed = parse_simple_option(out_path, rows)
    assert parsed.lines[0].quantity == 2
    assert parsed.lines[0].product_name_raw == "충전식 전동 분무기 BMC 18L"


def test_unknown_company_raises_instead_of_guessing(project_root, tmp_path):
    templates = load_output_templates(project_root / "config" / "output_templates.json")
    try:
        write_supplier_order_file("존재하지않는업체", [], templates, tmp_path / "x.xlsx")
        assert False, "예외가 발생해야 한다"
    except ValueError as e:
        assert "존재하지않는업체" in str(e)
